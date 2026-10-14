#!/usr/bin/env python3
"""Lecta ölçüm aracı — inceleme paketi.

  python3 inceleme.py uret [--durum draft] [--out ../output/inceleme-paketi.xlsx]
      Kaynaktaki maddelerden okuma/inceleme tablosu üretir (TR öncelikli, EN yan sütunlarda).
  python3 inceleme.py oku ../output/inceleme-paketi.xlsx
      Doldurulmuş "Karar" ve "Yorum" sütunlarını listeler (YAML'a işlenmek üzere).
  python3 inceleme.py tablo [--out ../output/madde-rubric-tablosu.md]
      Tüm maddelerin TR rubric tablosunu markdown olarak üretir.
"""
import argparse, glob, os, sys
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'source')
ORDER = {'governance': 0, 'matrix': 1}
LAYERS = ['CON', 'HYG', 'IDE', 'BEH', 'TRA', 'ATT']


def items():
    out = []
    for f in sorted(glob.glob(os.path.join(SRC, '*.yaml'))):
        if os.path.basename(f) in ('meta.yaml', '_template.yaml'):
            continue
        out += yaml.safe_load(open(f, encoding='utf-8'))['items']
    key = lambda i: (ORDER[i['block']], LAYERS.index(i['layer']) if i['block'] == 'matrix' else 0, i['id'])
    return sorted(out, key=key)


def j(x, lang='tr'):
    return ' '.join(str(x[lang]).split()) if x else ''


def cmd_uret(a):
    its = [i for i in items() if not a.durum or i.get('status') == a.durum]
    wb = Workbook()
    ws = wb.active; ws.title = 'Maddeler'
    cols = [('ID', 13), ('Karar', 11), ('Yorum', 40), ('Durum', 9), ('QS', 5), ('CE', 6), ('SOV', 6),
            ('Başlık', 28), ('Amaç', 50), ('Soru', 40), ('Uygulanabilirlik (NA)', 30),
            ('0', 30), ('1', 45), ('2', 45), ('3', 40), ('Gözden geçirme (ay)', 8),
            ('Kanıt', 45), ('Eşleme', 30), ('Servisler', 14), ('Not (iç)', 40),
            ('Title (EN)', 28), ('Rubric 2 (EN)', 45)]
    ws.append([c for c, _ in cols])
    for n, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, n).column_letter].width = w
    for i in its:
        r = i['rubric']
        ri = i.get('review_interval')
        rub = [j(r[k]).replace('review_interval', f'{ri} ay' if ri else 'review_interval') for k in (0, 1, 2, 3)]
        ev = '\n'.join(f"[{e['type']}] {' '.join(e['tr'].split())}" for e in i['evidence'])
        mp = '\n'.join(f"{m['framework']} {m['ref']}" for m in i['mappings'])
        lv = i['level']
        ws.append([i['id'], '', '', i.get('status', ''), '✓' if i.get('quick_scan') else '',
                   lv.get('CE') or '—', lv.get('SOV') or '—',
                   j(i['title']), j(i['objective']), j(i['question']), j(i['applicability']),
                   *rub, ri or '', ev, mp, ', '.join(i.get('services', [])),
                   ' '.join(str(i.get('notes', '')).split()),
                   j(i['title'], 'en'), j(r[2], 'en')])
    hdr = PatternFill('solid', fgColor='1F3A5F'); you = PatternFill('solid', fgColor='FFF4CC')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        row[1].fill = you; row[2].fill = you
        row[0].font = Font(bold=True)
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type='list', formula1='"Onay,Düzelt,Sil,Tartışalım"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'B2:B{ws.max_row}')
    # hidden EN columns grouped
    ws.column_dimensions.group('U', 'V', hidden=True)

    h = wb.create_sheet('Nasıl okunur', 0)
    lines = [
        ('Lecta Sovereign AI Control Framework — madde inceleme paketi', True),
        (f'{len(its)} madde · kaynak: measurement-tool/source/*.yaml · tools/inceleme.py ile üretildi', False),
        ('', False),
        ('1. "Maddeler" sekmesinde her satır bir madde. Sarı iki sütun senin: Karar (Onay / Düzelt / Sil / Tartışalım) ve Yorum.', False),
        ('2. Okuma sırası önerisi: önce QS ✓ olanlar (filtre: QS), sonra kalanlar. G maddeleri üstte, matris altta.', False),
        ('3. Her madde için bakılacak üç şey: (a) Amaç tek bir gereksinim mi? (b) Rubric 1 ile 2 arasındaki çizgiyi başka bir değerlendirici aynı yere çizer mi? (c) Seviye (CE/SOV) doğru mu?', False),
        ('4. "Düzelt" dediğinde yorumda ne değişecek, kısaca yaz; metni yeniden yazman gerekmez.', False),
        ('5. EN metinler U–V sütunlarında gizli (grup "+" ile açılır). TR ile anlam farkı görürsen yorumda belirt.', False),
        ('6. Bitirince dosyayı aynı yerde kaydet ve Claude\'a "inceleme bitti" de: kararlar okunur, YAML güncellenir, onaylananlar approved olur.', False),
    ]
    for t, b in lines:
        h.append([t]); h.cell(h.max_row, 1).font = Font(bold=b, size=13 if b else 11)
    h.column_dimensions['A'].width = 140
    wb.save(a.out)
    print(f'yazıldı: {a.out} ({len(its)} madde)')


def cmd_oku(a):
    ws = load_workbook(a.xlsx)['Maddeler']
    hdr = [c.value for c in ws[1]]
    ix = {k: hdr.index(k) for k in ('ID', 'Karar', 'Yorum')}
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        k, y = row[ix['Karar']], row[ix['Yorum']]
        if k or y:
            n += 1
            print(f"{row[ix['ID']]}\t{k or '-'}\t{y or ''}")
    print(f'# {n} madde işaretli', file=sys.stderr)


def cmd_tablo(a):
    def c(x):
        return ' '.join(str(x).split()).replace('|', '/')
    its = [i for i in items() if i.get('status') != 'deprecated']
    L = ['# Madde ve rubric tablosu', '',
         f'*Kaynak: `source/*.yaml` · `tools/inceleme.py tablo` ile üretildi — elle düzenlemeyin · {len(its)} madde*', '',
         'CE / SOV: maddenin o profildeki seviyesi (— = o profilde uygulanmaz) · QS: Quick Scan kesitinde · "… ay içinde": gözden geçirme aralığı', '',
         '| ID | Madde | CE | SOV | QS | 0 | 1 | 2 | 3 |', '| --- | --- | --- | --- | --- | --- | --- | --- | --- |']
    for i in its:
        ri = i.get('review_interval')
        r = [c(i['rubric'][k]['tr']).replace('review_interval', f'{ri} ay') for k in (0, 1, 2, 3)]
        lv = i['level']
        L.append(f"| {i['id']} | {c(i['title']['tr'])} | {lv.get('CE') or '—'} | {lv.get('SOV') or '—'} | "
                 f"{'✓' if i.get('quick_scan') else ''} | " + ' | '.join(r) + ' |')
    with open(a.out, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(L) + '\n')
    print(f'yazıldı: {a.out} ({len(its)} madde)')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest='cmd', required=True)
    u = sp.add_parser('uret'); u.add_argument('--durum', default=None)
    u.add_argument('--out', default=os.path.join(HERE, '..', 'output', 'inceleme-paketi.xlsx'))
    o = sp.add_parser('oku'); o.add_argument('xlsx')
    t = sp.add_parser('tablo'); t.add_argument('--out', default=os.path.join(HERE, '..', 'output', 'madde-rubric-tablosu.md'))
    a = ap.parse_args()
    {'uret': cmd_uret, 'oku': cmd_oku, 'tablo': cmd_tablo}[a.cmd](a)


if __name__ == '__main__':
    main()
